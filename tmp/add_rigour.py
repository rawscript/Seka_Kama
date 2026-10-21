from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PATH='Seka_Kama/docs/Seka Kama System Documentation - Final Draft.docx'
OUT='Seka_Kama/docs/Seka Kama System Documentation - Research Grade Draft.docx'
doc=Document(PATH)

def p_before(t,text='',style='Normal'):
    p=t.insert_paragraph_before(text); p.style=style; return p
def table_before(t,data):
    tab=doc.add_table(rows=len(data),cols=len(data[0])); tab.style='Table Grid'
    for i,row in enumerate(data):
        for j,val in enumerate(row):
            c=tab.cell(i,j); c.text=str(val); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs: r.font.name='Aptos'; r.font.size=Pt(8.2)
            if i==0:
                shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'5B6573'); c._tc.get_or_add_tcPr().append(shd)
                for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    t._p.addprevious(tab._tbl)

t=next(p for p in doc.paragraphs if p.text.startswith('CHAPTER FIVE: CONCLUSIONS'))
p_before(t,'4.6 Reproducible map implementation specification','Heading 2')
p_before(t,'The map workspace is implemented as a composition of named libraries, rather than as a single mapping package. This distinction matters because each library owns a different part of the interaction model. The following specification is taken from web-app/seka_kama/package.json and the component files in web-app/seka_kama/components. A library is listed as an implementation dependency only where the repository manifest records it; the document does not infer capabilities from a package name alone.','Normal')
p_before(t,'Table 4.7 Exact map and visualisation dependencies','Caption')
table_before(t,[
 ['Library and version','Integration point','Documented responsibility','Evidence source'],
 ['maplibre-gl ^5.24.0','Map canvas and style runtime','Renders the interactive basemap, camera, sources and map layers.','web-app/seka_kama/package.json; components/KeplerMap.tsx and SekaMap.tsx'],
 ['react-map-gl ^8.1.1','React map wrapper','Connects React state and lifecycle to the map instance used by the workspace.','web-app/seka_kama/package.json; map components'],
 ['@mapbox/mapbox-gl-draw ^1.5.1','Geometry editing control','Captures the user-defined area of interest as GeoJSON for scenario requests.','web-app/seka_kama/package.json; scenario UI components'],
 ['deck.gl 9.2.11','Analytical rendering layer','Provides GPU-accelerated layers for dense grid-cell and spatial overlays.','web-app/seka_kama/package.json; deck.gl overrides'],
 ['kepler.gl ^3.3.0-alpha.0','Exploratory geospatial view','Supports exploratory layer inspection and map-based analysis where enabled by the application.','web-app/seka_kama/package.json'],
 ['recharts ^2.15.3','Chart components','Renders time-series, comparison and summary charts alongside the map.','web-app/seka_kama/package.json; TrendChart.tsx'],
 ['lucide-react ^0.473.0','Interface icons','Provides consistent icons for controls and navigation; it is not a spatial analysis library.','web-app/seka_kama/package.json'],
])
p_before(t,'The map-to-model contract is explicit. The draw control produces GeoJSON; the client sends the geometry to the FastAPI scenario route; the backend converts it to a spatial query, identifies intersecting grid cells and returns result GeoJSON for the map. Baseline, predicted and scenario layers must be styled separately so a reader can distinguish observed or stored values from model output. The implementation must also expose the source date, model version, feature units and missing-data state in the layer metadata or result panel.','Normal')
p_before(t,'4.7 Exact model feature inventory','Heading 2')
p_before(t,'The prediction service loads 43 feature names from backend/models/feature_names.pkl. These names are the authoritative input order for the scaler and XGBoost model. The repository does not provide a complete scientific definition or unit for every feature; therefore the dictionary below records the name and family and marks the unit as requiring confirmation where the source schema does not define it. This prevents a software label from being presented as an ecological measurement without evidence.','Normal')
p_before(t,'Table 4.8 Authoritative model feature inventory','Caption')
features='longterm_slope_mean,longterm_slope_std,all_skew_mean,all_skew_std,all_mean_mean,all_mean_std,all_kurtosis_mean,all_kurtosis_std,all_median_mean,all_median_std,all_variance_mean,all_variance_std,licorr_slope_mean,licorr_slope_std,longterm_intercept_mean,longterm_intercept_std,longterm_r2_mean,longterm_r2_std,pop2018_mean,pop2018_std,primary_acf_mean,primary_acf_std,primary_prominence_mean,primary_prominence_std,secondary_acf_mean,secondary_acf_std,secondary_prominence_mean,secondary_prominence_std,ann_amp_mean,ann_amp_std,ann_cv_mean,ann_cv_std,ann_peak_month_mean,ann_peak_month_std,ann_trend_mean,ann_trend_std,ann_mean_mean,ann_mean_std,density_code,hist_lag1,hist_lag2,cheetah_abundance,dist_to_protected_km'.split(',')
rows=[['Index','Feature name','Feature family','Unit / definition status']]
for i,f in enumerate(features,1):
    if f.startswith('longterm_') or f.startswith('licorr_'): fam='Long-term / correlation trend'
    elif f.startswith('all_'): fam='Night-light distribution statistic'
    elif f.startswith('pop') or f.startswith('primary_') or f.startswith('secondary_'): fam='Population / temporal continuity proxy'
    elif f.startswith('ann_'): fam='Annual seasonality statistic'
    elif f.startswith('hist_'): fam='Historical lag feature'
    elif f=='density_code': fam='Encoded density feature'
    elif f=='cheetah_abundance': fam='Ecological proxy'
    else: fam='Spatial distance feature'
    unit='km for distance' if f=='dist_to_protected_km' else 'Unit requires confirmation from data dictionary / source product'
    rows.append([i,f,fam,unit])
table_before(t,rows)
p_before(t,'4.8 Methodological limits and unresolved definitions','Heading 2')
p_before(t,'Three items still require project-owner confirmation before publication as empirical results: the ecological units and source products for every model feature; the calibration and validation dataset for the 43-feature model bundle; and the exact production map tile provider and licensing terms. The software repository establishes the names and interfaces, but it does not by itself establish ecological validity. Those items should be completed in the data dictionary and methods chapter before the work is submitted as a research paper.','Normal')

doc.save(OUT)
print(OUT)
