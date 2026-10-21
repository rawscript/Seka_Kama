from docx import Document

SRC='Seka_Kama/docs/Seka Kama System Documentation - Explicit Map Libraries.docx'
OUT='Seka_Kama/docs/Seka Kama System Documentation - Fully Explicit.docx'
doc=Document(SRC)
repls={
 'A spatial scenario workflow is justified since conservation and land-use decisions are place-specific involving various interacting indicators. A digital-twin approach promises to connect reference layers, modelling and map-based review. The use of web workspace and geospatial services will be under':'A spatial scenario workflow is justified since conservation and land-use decisions are place-specific and involve interacting indicators. The digital-twin approach connects reference layers, modelling and review through the Next.js/React workspace, MapLibre GL JS, Mapbox tiles and the FastAPI geospatial services. The use of this web workspace and these geospatial services will be under',
 'The map-first workspace will combine the map with a scenario panel':'The Next.js/React workspace will combine the MapLibre GL JS and Mapbox map canvas with a scenario panel',
 'The presentation layer will provide maps and forms.':'The presentation layer will provide the MapLibre GL JS/Mapbox canvas, Deck.gl analytical layers and forms.',
 'The design will connect a map-based workflow to spatial queries':'The design will connect a MapLibre GL JS/Mapbox workflow to spatial queries',
 'The map is the primary working surface for Seka Kama.':'The MapLibre GL JS canvas, configured with Mapbox styles and tiles, is the primary working surface for Seka Kama.',
 'The map workflow should remain understandable':'The MapLibre GL JS layer workflow should remain understandable',
 'The map workspace is implemented as a composition of named libraries':'The Next.js/React map workspace is implemented as a composition of named libraries',
 'MapLibre handles the base map, Draw captures the area of interest, Deck.gl renders analytical layers, and the chart components explain the selected results.':'MapLibre GL JS renders the Mapbox basemap, @mapbox/mapbox-gl-draw captures the area of interest, Deck.gl renders analytical layers, and Recharts explains the selected results.',
 'returns result GeoJSON for the map.':'returns result GeoJSON for rendering by MapLibre GL JS and Deck.gl.',
 'Figure 2.2 Proposed map workspace reference [ Page.  ]':'Figure 2.2 Proposed Next.js/React and MapLibre GL JS workspace reference [ Page.  ]',
 'Figure 2.2 Proposed map-workspace reference; screen content will be adapted to the approved study workflow.':'Figure 2.2 Proposed Next.js/React and MapLibre GL JS workspace reference; screen content will be adapted to the approved study workflow.',
 'Table 4.1 Map and interface tools':'Table 4.1 MapLibre GL JS, Mapbox and interface tools',
}
for p in doc.paragraphs:
    for a,b in repls.items():
        if a in p.text: p.text=p.text.replace(a,b)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for a,b in repls.items():
                    if a in p.text: p.text=p.text.replace(a,b)
doc.save(OUT)
print(OUT)
