from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC='Seka_Kama/docs/Seka Kama System Documentation - Submission Draft.docx'
OUT='Seka_Kama/docs/Seka Kama System Documentation - Explicit Map Libraries.docx'
doc=Document(SRC)
target=None
for t in doc.tables:
    if t.rows and t.rows[0].cells[0].text.strip()=='Resource':
        target=t; break
if target is None: raise RuntimeError('Map table not found')
rows=[
 ['Library / package','Version in project','Explicit use in Seka Kama','Integration point'],
 ['maplibre-gl','^5.24.0','Renders the interactive map canvas, basemap style, camera, sources and visible map layers.','SekaMap.tsx and KeplerMap.tsx'],
 ['react-map-gl','^8.1.1','Provides the React component and event wrapper around the map instance.','Map workspace components'],
 ['@mapbox/mapbox-gl-draw','^1.5.1','Lets the user draw and edit the scenario area of interest; outputs GeoJSON geometry.','Scenario drawing control'],
 ['deck.gl','9.2.11','Renders analytical grid-cell and spatial overlay layers using GPU-accelerated WebGL layers.','Analytical map layers'],
 ['kepler.gl','^3.3.0-alpha.0','Provides exploratory geospatial layer inspection and configuration where enabled.','Exploratory map view'],
 ['recharts','^2.15.3','Renders the trend, comparison and summary charts associated with the selected map data.','TrendChart.tsx and result panels'],
 ['lucide-react','^0.473.0','Provides interface icons for map controls, panels and navigation; it does not render spatial data.','Shared interface components'],
]
new=doc.add_table(rows=len(rows), cols=4); new.style='Table Grid'; target._tbl.addprevious(new._tbl)
for i,row in enumerate(rows):
    for j,val in enumerate(row):
        c=new.cell(i,j); c.text=val; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(2)
            for r in p.runs: r.font.name='Aptos'; r.font.size=Pt(8.5)
        if i==0:
            shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'5B6573'); c._tc.get_or_add_tcPr().append(shd)
            for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
target._element.getparent().remove(target._element)
doc.save(OUT)
print(OUT)
