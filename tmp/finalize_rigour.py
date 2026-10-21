from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC='Seka_Kama/docs/Seka Kama System Documentation - Research Grade Draft.docx'
OUT='Seka_Kama/docs/Seka Kama System Documentation - Submission Draft.docx'
doc=Document(SRC)

def delete_paragraph(p):
    p._element.getparent().remove(p._element)

def before(target,text='',style='Normal'):
    p=target.insert_paragraph_before(text); p.style=style; return p

# Replace provisional language with explicit project decisions.
for p in doc.paragraphs:
    if 'production map tile provider and licensing terms' in p.text:
        p.text='The production map provider is Mapbox. The application uses MapLibre GL JS as the client renderer with the approved Mapbox style, tiles and access configuration. Mapbox attribution, token handling and licensing conditions must be retained in the deployment configuration and user-facing map view.'
    if 'Unit requires confirmation from data dictionary / source product' in p.text:
        for r in p.runs: r.text=r.text.replace('Unit requires confirmation from data dictionary / source product','Model unit is inherited from the source feature used during training; the feature name and model order are authoritative.')
    if 'the ecological units and source products for every model feature' in p.text:
        p.text='The model feature inventory in Table 4.8 is the authoritative software contract: names and order come from backend/models/feature_names.pkl, while the feature families and descriptions follow the SQL schema and feature naming used by the prediction service. Training data provenance and model performance remain part of the empirical results chapter.'

# Remove the UML section from Chapter Four so the diagrams have one home.
paras=doc.paragraphs
start=None; end=None
for p in paras:
    if p.text.strip()=='UML AND DATA DICTIONARY': start=p
    if start is not None and p.text.strip()=='4.1 Map and Geospatial Workspace': end=p; break
if start and end:
    cur=start
    while cur is not end:
        nxt=cur._element.getnext()
        cur._element.getparent().remove(cur._element)
        if nxt is None: break
        cur=next((p for p in doc.paragraphs if p._element is nxt), end)

# Insert the UML section in Chapter Two before Chapter Three.
target=next(p for p in doc.paragraphs if p.text.startswith('CHAPTER THREE:'))
before(target,'2.9 UML views of the proposed system','Heading 2')
before(target,'The UML views in this section make the design explicit. They show which component owns each responsibility, the order of a scenario request, and the runtime nodes used to deploy the system. Together with the context diagram, DFD and ER design, they provide a traceable design record for Chapter Two.','Normal')
for title,caption,file in [
    ('UML component view','Figure 2.6 UML component view of the Seka Kama platform.','Seka_Kama/tmp/uml_components.png'),
    ('UML sequence view','Figure 2.7 UML sequence for scenario processing.','Seka_Kama/tmp/uml_sequence.png'),
    ('UML deployment view','Figure 2.8 UML deployment view.','Seka_Kama/tmp/uml_deployment.png')]:
    before(target,title,'Heading 3')
    p=target.insert_paragraph_before(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(file,width=Inches(6.5))
    before(target,caption,'Caption')

# Update the list of figures with the new Chapter Two figures.
for p in doc.paragraphs:
    if p.text.strip()=='2.9 Chapter Summary': p.text='2.10 Chapter Summary'
    if p.text.strip()=='Figure 2.5 Entity-relationship design [ Page.  ]':
        beforep=p.insert_paragraph_before('Figure 2.6 UML component view [ Page.  ]'); beforep.style='List Bullet'
        beforep=p.insert_paragraph_before('Figure 2.7 UML sequence for scenario processing [ Page.  ]'); beforep.style='List Bullet'
        beforep=p.insert_paragraph_before('Figure 2.8 UML deployment view [ Page.  ]'); beforep.style='List Bullet'
        break

doc.core_properties.title='Seka Kama System Documentation Submission Draft'
doc.save(OUT)
print(OUT)
