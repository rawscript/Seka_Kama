from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC='Seka_Kama/docs/Seka Kama System Documentation - Fully Explicit.docx'
OUT='Seka_Kama/docs/Seka Kama System Documentation - Complete Tools.docx'
doc=Document(SRC)
target=None
for t in doc.tables:
    if t.rows and t.rows[0].cells[0].text.strip()=='Library / package': target=t; break
if target is None: raise RuntimeError('Explicit tool table not found')
rows=[
 ['Tool / package','Version or source','Explicit responsibility','Integration point'],
 ['Next.js','16.2.6','Application shell, routing, server/client rendering and production web build.','web-app/seka_kama/app and next.config.ts'],
 ['React','19.2.4','Component model and client-side interaction state for the web workspace.','web-app/seka_kama/components'],
 ['TypeScript','project devDependency','Typed client services, components and API contracts.','web-app/seka_kama/*.ts and *.tsx'],
 ['maplibre-gl','^5.24.0','Renders the Mapbox basemap, camera, sources and visible map layers.','SekaMap.tsx and KeplerMap.tsx'],
 ['react-map-gl','^8.1.1','Connects the MapLibre GL JS map instance to React components and events.','Map workspace components'],
 ['@mapbox/mapbox-gl-draw','^1.5.1','Captures and edits scenario geometry as GeoJSON.','Scenario drawing control'],
 ['deck.gl','9.2.11','Renders analytical grid-cell and spatial overlay layers.','Analytical map layers'],
 ['kepler.gl','^3.3.0-alpha.0','Supports exploratory geospatial layer inspection where enabled.','Exploratory map view'],
 ['recharts','^2.15.3','Renders trend, comparison and summary charts.','TrendChart.tsx and result panels'],
 ['lucide-react','^0.473.0','Provides interface icons for controls and navigation.','Shared interface components'],
 ['Python / FastAPI','backend requirements','Provides the HTTP API, validation, authentication and orchestration services.','backend/api, backend/core and backend/main.py'],
 ['PostGIS / PostgreSQL','Docker image postgis/postgis:16-3.4','Stores grid cells, protected areas, users, scenarios, history and audit records with spatial indexes.','backend/sql and docker-compose.yml'],
 ['Supabase client','backend requirements','Provides the application data-access connection to the hosted PostgreSQL/PostGIS store.','backend/core/database.py and services'],
 ['XGBoost model','backend/models/sekanet_xgboost_shp.pkl','Produces the model prediction from the ordered feature vector.','backend/services/prediction_service.py'],
 ['StandardScaler','backend/models/sekanet_scaler_shp.pkl','Transforms the model feature vector before XGBoost inference.','backend/services/prediction_service.py'],
 ['NASA POWER / CHIRPS','backend integration scripts','Supplies rainfall and related environmental enrichment when enabled.','backend/scripts/fetch_rainfall_*.py'],
 ['GBIF API','backend/services/ecological_data_service.py','Supplies species-occurrence information used for ecological enrichment.','Ecological data service'],
 ['NVIDIA NIM / OpenAI-compatible client','backend/services/llm_service.py','Generates supplementary narrative explanations with a rule-based fallback.','LLM service'],
 ['Docker Compose','docker-compose.yml','Runs the PostGIS, FastAPI and Next.js services as separate containers.','Deployment configuration'],
 ['JWT, bcrypt and API keys','backend authentication modules','Protects authenticated routes and supports programmatic access.','backend/api/auth_routes.py and key_routes.py'],
]
new=doc.add_table(rows=len(rows),cols=4); new.style='Table Grid'; target._tbl.addprevious(new._tbl)
for i,row in enumerate(rows):
    for j,val in enumerate(row):
        c=new.cell(i,j); c.text=val; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(2)
            for r in p.runs: r.font.name='Aptos'; r.font.size=Pt(8.2)
        if i==0:
            shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'5B6573'); c._tc.get_or_add_tcPr().append(shd)
            for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
target._element.getparent().remove(target._element)
doc.save(OUT)
print(OUT)
