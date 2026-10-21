from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC=r'C:\Users\LENOVO T14\Downloads\Seka Kama System Documentation (2).docx'
OUT='Seka_Kama/docs/Seka Kama System Documentation - Existing Systems Review.docx'
doc=Document(SRC)

def before(t,text='',style='Normal'):
    p=t.insert_paragraph_before(text); p.style=style; return p

def add_table_before(t,data):
    tab=doc.add_table(rows=len(data),cols=len(data[0])); tab.style='Table Grid'
    for i,row in enumerate(data):
        for j,val in enumerate(row):
            c=tab.cell(i,j); c.text=str(val); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs: r.font.name='Aptos'; r.font.size=Pt(8.5)
            if i==0:
                shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'5B6573'); c._tc.get_or_add_tcPr().append(shd)
                for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    t._p.addprevious(tab._tbl)

target=next(p for p in doc.paragraphs if p.text.strip()=='1.3 System Objectives')
before(target,'1.3 Existing Systems and Comparative Gap','Heading 2')
before(target,'Existing conservation systems demonstrate that the problem is technically tractable, but they address different operational needs. The comparison below is based on the published capabilities of each system and is used to define the gap for this project. It does not claim that Seka Kama replaces any of them.','Normal')
before(target,'SMART is a conservation monitoring and reporting platform with mobile, desktop and cloud components. It supports structured patrol and observation data, analysis and reporting across a conservation area. Its primary strength is operational monitoring and reporting. The project requirement that remains open is a scenario workflow that modifies spatial predictor values and compares modelled outcomes against a baseline.','Normal')
before(target,'EarthRanger is a real-time conservation platform that integrates GPS collars, camera traps, patrol reports and remote sensors into a unified map view. It supports wildlife, personnel and vehicle tracking, offline field collection and configurable alerts. Its primary strength is operational awareness and response. This project addresses a different analytical need: scenario-based comparison of habitat and human-pressure variables for conservation planning.','Normal')
before(target,'ArcGIS provides mature spatial analysis, suitability modelling and conservation-planning workflows. It is a general GIS environment rather than a single lion-conservation scenario service. Its strength is broad geoprocessing and professional cartography. Seka Kama narrows the workflow to an authenticated web workspace, an approved feature vector, an XGBoost prediction service and a reproducible scenario record.','Normal')
before(target,'Table 1.1 Existing-system comparison','Caption')
add_table_before(target,[
 ['System','Primary purpose','Relevant capabilities','Limitation relative to this study','Source'],
 ['SMART','Conservation monitoring and reporting','Mobile, desktop and cloud components; structured patrol and observation data; analysis and reporting.','Not documented here as a purpose-built workflow for changing model predictors and comparing predicted lion outcomes.','SMART 7 overview, World Bank: https://thedocs.worldbank.org/en/doc/43c89c4a946e031980277302e3072960-0320052022/original/poaching-spatial-management-and-reporting-tool-smart.pdf'],
 ['EarthRanger','Real-time wildlife and protected-area operations','Integrates GPS collars, camera traps, patrol reports and sensors; unified map; offline collection; alerts.','Designed around monitoring and response rather than the specific baseline-versus-scenario modelling workflow proposed here.','EarthRanger Support: https://support.earthranger.com/en_US/what-is-earthranger-and-how-does-it-work'],
 ['ArcGIS conservation workflows','General GIS analysis and conservation planning','Suitability modelling, spatial analysis, prioritisation and mapping.','Broad platform; the proposed study defines a narrower, web-based, model-linked scenario workflow for the Greater Mara.','Esri Conservation Planning: https://www.esri.com/en-us/industries/conservation/focus-areas/planning'],
 ['Seka Kama proposed platform','Scenario-based lion conservation planning','Mapbox/MapLibre GL JS workspace; GeoJSON area selection; PostGIS spatial query; 43-feature model vector; XGBoost prediction; baseline comparison; scenario history.','Requires project evaluation before claims about accuracy, operational benefit or ecological validity can be made.','This study design and implementation repository.'],
])
before(target,'The comparative gap is therefore specific. Existing systems provide strong monitoring, real-time operations or general GIS analysis. Seka Kama is designed to connect a selected geography to an approved 43-feature model vector, apply permitted scenario changes, run a prediction, return baseline-versus-scenario differences and retain the request for review. The evaluation must test whether this proposed workflow is useful and valid; the comparison does not establish that it is superior.','Normal')

# Shift later section numbers to make room for 1.3.
for p in doc.paragraphs:
    repl={'1.3 System Objectives':'1.4 System Objectives','1.3.1 General Objective':'1.4.1 General Objective','1.3.2 Specific Objectives':'1.4.2 Specific Objectives','1.3.3 Research Questions':'1.4.3 Research Questions','1.3.4 Working Hypothesis':'1.4.4 Working Hypothesis','1.4 Functional Requirements':'1.5 Functional Requirements','1.5 Non-Functional Requirements':'1.6 Non-Functional Requirements','1.6 Scope and Delimitation':'1.7 Scope and Delimitation','1.7 Tools and Resources':'1.8 Tools and Resources','1.8 Project Schedule Breakdown':'1.9 Project Schedule Breakdown'}
    if p.text.strip() in repl: p.text=repl[p.text.strip()]

# Add authoritative sources to references.
refs=next(p for p in doc.paragraphs if p.text.strip()=='REFERENCES')
before(refs,'Allen Institute for AI. (2026). What is EarthRanger and how does it work? EarthRanger Support. https://support.earthranger.com/en_US/what-is-earthranger-and-how-does-it-work','Normal')
before(refs,'Esri. (n.d.). Conservation planning and prioritization. https://www.esri.com/en-us/industries/conservation/focus-areas/planning','Normal')
before(refs,'World Bank. (2022). Spatial Management and Reporting Tool (SMART). https://thedocs.worldbank.org/en/doc/43c89c4a946e031980277302e3072960-0320052022/original/poaching-spatial-management-and-reporting-tool-smart.pdf','Normal')

doc.save(OUT)
print(OUT)
