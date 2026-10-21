from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PATH='Seka_Kama/docs/Seka Kama System Documentation - Revised.docx'
OUT='Seka_Kama/docs/Seka Kama System Documentation - Publication Draft.docx'
doc=Document(PATH)

def insert_p_before(target, text='', style=None):
    p=target.insert_paragraph_before(text)
    if style: p.style=style
    return p

def insert_table_before(target, rows, cols, data, widths=None):
    t=doc.add_table(rows=rows, cols=cols)
    t.style='Table Grid'
    for i,row in enumerate(data):
        for j,val in enumerate(row):
            c=t.cell(i,j); c.text=str(val); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs: r.font.name='Aptos'; r.font.size=Pt(9)
            if i==0:
                shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'5B6573'); c._tc.get_or_add_tcPr().append(shd)
                for p in c.paragraphs:
                    for r in p.runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    target._p.addprevious(t._tbl)
    return t

target=next(p for p in doc.paragraphs if p.text.strip()=='CHAPTER FOUR: CONCLUSIONS AND RECOMMENDATIONS')

insert_p_before(target,'CHAPTER FOUR: TECHNICAL IMPLEMENTATION AND OPERATIONAL DOCUMENTATION','Heading 1')
insert_p_before(target,'This chapter records the implementation tools and operational controls that make the proposed platform reproducible. The descriptions are tied to the current project structure and distinguish implemented components from planned integrations.','Normal')
insert_p_before(target,'4.1 Map and Geospatial Workspace','Heading 2')
insert_p_before(target,'The map is the primary working surface for Seka Kama. The web client is built with Next.js and React. MapLibre GL JS provides the map renderer and camera controls, while react-map-gl supplies the React integration. Mapbox Draw supports user-defined polygons and other geometry inputs. Deck.gl is used for GPU-assisted rendering of dense grid-cell layers and derived spatial overlays. Kepler.gl is retained for exploratory visual analysis where an analyst needs rapid layer inspection. Recharts provides compact trend and comparison charts beside the map. These tools serve different purposes: MapLibre handles the base map, Draw captures the area of interest, Deck.gl renders analytical layers, and the chart components explain the selected results.','Normal')
insert_p_before(target,'The map workflow should remain understandable when a layer is hidden or unavailable. Each layer therefore has a name, source date, unit, visible range and provenance note. Scenario changes are drawn in a separate style from baseline observations, and the legend states whether a value is observed, derived or predicted. Sensitive wildlife locations are generalized or withheld according to the user role.','Normal')
insert_p_before(target,'Table 4.1 Map and interface tools','Caption')
insert_table_before(target,5,4,[
 ['Tool','Role in the platform','Primary output','Control to document'],
 ['Next.js and React','Application shell, routing and stateful workspace','Browser interface','Version, build command and environment variables'],
 ['MapLibre GL JS and react-map-gl','Base map, navigation and layer display','Interactive map canvas','Coordinate reference system and tile source'],
 ['Mapbox Mapbox Draw','Area-of-interest geometry and editing','GeoJSON geometry','Geometry validation and permitted bounds'],
 ['Deck.gl, Kepler.gl and Recharts','Analytical layers, exploratory inspection and charts','Grid layers, summaries and trends','Layer metadata, scales and units'],
],None)
insert_p_before(target,'4.2 Spatial Data Pipeline','Heading 2')
insert_p_before(target,'Spatial data enters the system through controlled ingestion scripts and database loaders. The current repository includes scripts for rainfall, historical trends, protected areas, land cover, NDVI and grid-cell imports. The target store is PostGIS, accessed through the Supabase client in the backend services. During ingestion, each layer is checked for geometry validity, coordinate reference system, temporal coverage, field names, units and source provenance. A layer is not made available to scenario analysis until these checks pass.','Normal')
insert_p_before(target,'The analytical unit is the grid cell. A baseline request can be filtered by management unit, year or bounding box. A scenario request supplies GeoJSON geometry; the spatial service converts it to a database query, identifies intersecting cells and returns a copy of the baseline features for processing. The baseline record remains read-only throughout this operation.','Normal')
insert_p_before(target,'4.3 Prediction and Scenario Services','Heading 2')
insert_p_before(target,'The backend is a Python FastAPI service. Authentication and API routing are separated from the domain services. The spatial service retrieves cells and protected areas; the ecological data service enriches cells with approved NASA POWER and GBIF values when those integrations are enabled; the prediction service loads the XGBoost model, StandardScaler and feature-name list; and the scenario processor applies only permitted feature modifications before inference. The response builder returns baseline and predicted totals, the difference between them, affected management units, the model version and a GeoJSON representation for the map.','Normal')
insert_p_before(target,'A narrative service may provide an ecological explanation through the configured NVIDIA NIM endpoint. The narrative is supplementary to the numeric result. If the language service is unavailable, the workflow returns a rule-based explanation and records that fallback. The system does not present generated narrative as an observed field event.','Normal')
insert_p_before(target,'Table 4.2 Core implementation components','Caption')
insert_table_before(target,7,3,[
 ['Component','Repository location','Responsibility'],
 ['FastAPI application','backend/api and backend/main.py','Routes, validation, authentication and response handling'],
 ['Spatial service','backend/services/spatial_service.py','Geometry conversion, intersections and spatial queries'],
 ['Prediction service','backend/services/prediction_service.py','Feature preparation, scaling and XGBoost inference'],
 ['Ecological data service','backend/services/ecological_data_service.py','NASA POWER and GBIF enrichment workflows'],
 ['Supabase/PostGIS access','backend/core/database.py and SQL scripts','Persistent spatial, user and audit records'],
 ['Next.js client','web-app/seka_kama/app and components','Map workspace, scenario controls and result views'],
],None)
insert_p_before(target,'4.4 API and Data Contracts','Heading 2')
insert_p_before(target,'The public contract is expressed through versioned HTTP endpoints. Health and status checks report database and model availability. Baseline endpoints return GeoJSON FeatureCollections with cell identifiers, management units, model features and dates. The scenario endpoint accepts a geometry, permitted feature modifications, optional management units, an optional user query and a simulation horizon. The response includes a scenario identifier, baseline and predicted totals, delta values, affected units, model metadata and visualization GeoJSON. Validation errors use a clear status code and message; a failed external enrichment is marked as a degraded enrichment rather than silently treated as a measured value.','Normal')
insert_p_before(target,'4.5 Deployment and Configuration','Heading 2')
insert_p_before(target,'The supplied Docker Compose configuration separates the PostGIS database, FastAPI backend and Next.js frontend. The database persists its data in a named volume. The backend receives the database URL, service key and model path through environment variables, and the frontend is exposed on its own service port. Production configuration must replace development secrets, restrict CORS origins, enable HTTPS at the edge and configure health checks before the system is exposed to users.','Normal')
insert_p_before(target,'Table 4.3 Deployment controls','Caption')
insert_table_before(target,6,3,[
 ['Area','Required control','Evidence to retain'],
 ['Secrets','Keep database keys, JWT secrets and external API keys outside source code','Redacted environment template and deployment record'],
 ['Database','Use PostGIS indexes, backups and a documented migration sequence','Migration log and backup test'],
 ['Model files','Pin model, scaler and feature-name versions together','Model manifest and checksum'],
 ['Web services','Run health checks and use HTTPS with restricted CORS','Health-check output and configuration review'],
 ['Audit and recovery','Record privileged actions and define rollback ownership','Audit sample and recovery runbook'],
],None)

# Renumber later chapter labels from 4.x to 5.x and appendices untouched.
for p in doc.paragraphs:
    if p.text.startswith('CHAPTER FOUR: CONCLUSIONS'):
        p.text=p.text.replace('CHAPTER FOUR','CHAPTER FIVE')
    elif p.text.startswith('4.1 Conclusion'):
        p.text=p.text.replace('4.1','5.1')
    elif p.text.startswith('4.2 Recommendations'):
        p.text=p.text.replace('4.2','5.2')
    elif p.text.startswith('4.3 Areas for Further Work'):
        p.text=p.text.replace('4.3','5.3')

doc.core_properties.title='Seka Kama System Documentation Publication Draft'
doc.save(OUT)
print(OUT)
