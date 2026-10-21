from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

PATH='Seka_Kama/docs/Seka Kama System Documentation - Publication Draft.docx'
OUT='Seka_Kama/docs/Seka Kama System Documentation - Final Draft.docx'
ASSET='Seka_Kama/tmp'

def font(size,bold=False):
    try: return ImageFont.truetype('C:/Windows/Fonts/arial.ttf',size=size)
    except: return ImageFont.load_default()

def diagram(path,title,boxes,arrows,size=(1800,1050)):
    im=Image.new('RGB',size,'white'); d=ImageDraw.Draw(im)
    d.text((50,30),title,fill='black',font=font(34,True))
    for name,(x,y,w,h) in boxes.items():
        d.rounded_rectangle((x,y,x+w,y+h),radius=16,outline='#334155',width=3,fill='#F3F4F6')
        lines=name.split('\n'); total=len(lines)*28
        for i,line in enumerate(lines):
            bb=d.textbbox((0,0),line,font=font(22,True)); tw=bb[2]-bb[0]
            d.text((x+(w-tw)/2,y+h/2-total/2+i*28),line,fill='#111827',font=font(22,True))
    for a,b,label in arrows:
        x1,y1,w1,h1=boxes[a]; x2,y2,w2,h2=boxes[b]
        sx=x1+w1; sy=y1+h1/2; ex=x2; ey=y2+h2/2
        if x2 < x1: sx=x1; ex=x2+w2
        d.line((sx,sy,ex,ey),fill='#475569',width=3)
        d.polygon([(ex,ey),(ex-12 if ex>sx else ex+12,ey-7),(ex-12 if ex>sx else ex+12,ey+7)],fill='#475569')
        if label: d.text(((sx+ex)/2-60,(sy+ey)/2-25),label,fill='#334155',font=font(18))
    im.save(path)

diagram(f'{ASSET}/uml_components.png','UML component view',{
 'Web client\nNext.js + React':(80,260,300,160),'Map workspace\nMapLibre + Deck.gl':(500,120,360,160),'FastAPI\nrouters + middleware':(500,420,360,160),'Domain services\nspatial / prediction / audit':(980,300,420,180),'PostGIS / Supabase\nspatial and history':(1500,150,250,160),'External data\nNASA / GBIF / NIM':(1500,500,250,160)},[
 ('Web client\nNext.js + React','Map workspace\nMapLibre + Deck.gl','renders'),('Web client\nNext.js + React','FastAPI\nrouters + middleware','HTTPS'),('FastAPI\nrouters + middleware','Domain services\nspatial / prediction / audit','calls'),('Domain services\nspatial / prediction / audit','PostGIS / Supabase\nspatial and history','queries'),('Domain services\nspatial / prediction / audit','External data\nNASA / GBIF / NIM','enriches')])

diagram(f'{ASSET}/uml_sequence.png','UML sequence for a scenario request',{
 'Analyst':(60,120,180,100),'Map client':(350,120,220,100),'API':(700,120,220,100),'Spatial service':(1050,120,240,100),'Prediction service':(1400,120,260,100),'Database':(1050,700,240,100)},[
 ('Analyst','Map client','draw area'),('Map client','API','POST /scenario'),('API','Spatial service','validate geometry'),('Spatial service','Database','intersect cells'),('API','Prediction service','apply changes'),('Prediction service','Database','save history'),('Prediction service','Map client','results'),('Map client','Analyst','map + summary')],(1800,980))

diagram(f'{ASSET}/uml_deployment.png','UML deployment view',{
 'Browser\nanalyst workstation':(80,350,300,160),'Frontend container\nNext.js :3000':(600,150,360,160),'Backend container\nFastAPI :8000':(600,550,360,160),'PostGIS volume\nPostgreSQL':(1300,150,330,160),'Model bundle\nXGBoost + scaler':(1300,550,330,160)},[
 ('Browser\nanalyst workstation','Frontend container\nNext.js :3000','HTTPS'),('Frontend container\nNext.js :3000','Backend container\nFastAPI :8000','JSON/GeoJSON'),('Backend container\nFastAPI :8000','PostGIS volume\nPostgreSQL','SQL/RPC'),('Backend container\nFastAPI :8000','Model bundle\nXGBoost + scaler','loads')])

doc=Document(PATH)
def p_before(t,text='',style=None):
    p=t.insert_paragraph_before(text); p.style=style or 'Normal'; return p
def table_before(t,data):
    tab=doc.add_table(rows=len(data),cols=len(data[0])); tab.style='Table Grid'
    for i,row in enumerate(data):
        for j,val in enumerate(row):
            c=tab.cell(i,j); c.text=str(val); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i==0:
                shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'5B6573'); c._tc.get_or_add_tcPr().append(shd)
                for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs: r.font.name='Aptos'; r.font.size=Pt(8.5)
    t._p.addprevious(tab._tbl)

t=next(p for p in doc.paragraphs if p.text.startswith('CHAPTER FOUR: TECHNICAL'))
p_before(t,'UML AND DATA DICTIONARY','Heading 1')
p_before(t,'This section complements the architecture description with diagrams and a consolidated data dictionary. The diagrams describe the current software boundary; the dictionary separates stored fields from calculated outputs and records the unit or interpretation needed for reproducible analysis.','Normal')
p_before(t,'UML component diagram','Heading 2'); p_before(t,'Figure 4.1 shows the principal software components and their dependencies.','Normal'); p=t.insert_paragraph_before(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(f'{ASSET}/uml_components.png',width=Inches(6.5)); p_before(t,'Figure 4.1 UML component view of the Seka Kama platform.','Caption')
p_before(t,'UML sequence diagram','Heading 2'); p_before(t,'Figure 4.2 follows the main scenario request from map interaction to spatial query, prediction, persistence and response.','Normal'); p=t.insert_paragraph_before(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(f'{ASSET}/uml_sequence.png',width=Inches(6.5)); p_before(t,'Figure 4.2 UML sequence for scenario processing.','Caption')
p_before(t,'UML deployment diagram','Heading 2'); p_before(t,'Figure 4.3 identifies the runtime nodes used by the Docker deployment and the model bundle loaded by the backend.','Normal'); p=t.insert_paragraph_before(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(f'{ASSET}/uml_deployment.png',width=Inches(6.5)); p_before(t,'Figure 4.3 UML deployment view.','Caption')
p_before(t,'Consolidated data dictionary','Heading 2')
p_before(t,'The project already contained fragments of a dictionary in the API documentation, model feature list and SQL bootstrap script. The following tables consolidate those fragments into a single reference for implementation, testing and review.','Normal')
p_before(t,'Table 4.4 Core spatial and ecological fields','Caption')
table_before(t,[
 ['Field','Type / unit','Meaning and use','Status'],
 ['cell_id','integer','Stable identifier for a grid-cell record.','Stored'],['geom','Polygon, EPSG:4326','Grid-cell boundary used for spatial intersection.','Stored'],['centroid','Point, EPSG:4326','Representative point used for bounding-box and enrichment queries.','Stored'],['management_unit','text','Conservancy or management area associated with the cell.','Stored'],['baseline_lion_density','float','Baseline density or model input recorded for the cell; interpretation must follow the approved model documentation.','Stored'],['year','integer','Observation or reference year for the cell record.','Stored'],['prey_density','float, source unit','Cached ecological proxy from approved prey data.','Optional cache'],['annual_rainfall_mm','float, mm','Annual rainfall used for enrichment or scenario analysis.','Optional cache'],['hwc_risk_score','float, 0-1','Derived human-wildlife conflict risk score.','Derived'],['dist_to_protected_km','float, km','Distance from the cell to the nearest protected area.','Derived / stored'],
])
p_before(t,'Table 4.5 Model feature families','Caption')
table_before(t,[
 ['Family','Representative fields','Interpretation'],['Night-light statistics','all_mean_mean, all_mean_std, all_skew_mean, all_kurtosis_mean, all_variance_mean','Summary statistics describing intensity and distribution of night-light observations.'],['Night-light trend','longterm_slope_mean, longterm_intercept_mean, longterm_r2_mean, licorr_slope_mean','Long-term level, slope and fit indicators.'],['Population and continuity','pop2018_mean, primary_acf_mean, secondary_acf_mean, prominence fields','Population proxy and temporal continuity features.'],['Seasonality','ann_amp_mean, ann_cv_mean, ann_peak_month_mean, ann_trend_mean','Annual amplitude, variation, peak timing and trend.'],['Ecological proxies','density_code, hist_lag1, hist_lag2, cheetah_abundance','Encoded or lagged ecological predictors used by the model.'],
])
p_before(t,'Table 4.6 User, scenario and audit entities','Caption')
table_before(t,[
 ['Entity','Key fields','Purpose'],['users','id, email, role, is_active, created_at','Identity, authorisation and account state.'],['api_keys','id, user_id, key_hash, is_active, last_used','Programmatic access without storing plaintext keys.'],['scenario_history','id, user_id, modified_features, request_data, scenario_geojson, baseline_total_lions, predicted_total_lions, delta_percent','Reproducible record of a scenario request and its result.'],['historical_stats','management_unit, year, population_estimate','Time-series values used in trends and comparison views.'],['audit_logs','user_id, action, resource_type, resource_id, details, created_at','Trace of privileged or material system actions.'],['protected_areas','site_name, designation, iucn_category, area_km2, geom','Protected-area boundaries and attributes.'],
])

doc.save(OUT)
print(OUT)
